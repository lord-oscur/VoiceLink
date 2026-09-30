"""Files attached to a message: text extraction, long-audio transcription and
conversion of stored attachments into OpenAI-style message content.

Attachment records (as stored in the conversation):
    {"id", "kind": "image" | "document" | "audio" | "video", "name",
     "data": data-URL (image), "text": extracted text / transcript,
     "frames": [data-URL, ...] (video), "duration": seconds, "truncated": bool}
"""
import io
import os

import numpy as np

MAX_DOC_CHARS = 60000
TEXT_EXT = {
    ".txt", ".md", ".markdown", ".csv", ".tsv", ".json", ".jsonl", ".xml", ".html", ".htm", ".log", ".ini",
    ".toml", ".yaml", ".yml", ".py", ".js", ".mjs", ".ts", ".tsx", ".jsx", ".java", ".c", ".h", ".cpp",
    ".hpp", ".cs", ".go", ".rs", ".rb", ".php", ".sh", ".bat", ".ps1", ".sql", ".css", ".scss", ".lua",
    ".kt", ".swift", ".r", ".m", ".tex", ".srt", ".vtt", ".rtf", ".cfg", ".conf", ".env",
}
DOC_EXT = TEXT_EXT | {".pdf", ".docx"}


class UnsupportedFile(ValueError):
    pass


def extract_text(name, data):
    """Returns (text, truncated) for a document file."""
    ext = os.path.splitext(name.lower())[1]
    if ext == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                raise UnsupportedFile("the PDF is password protected")
        pages = []
        for i, page in enumerate(reader.pages, 1):
            t = (page.extract_text() or "").strip()
            if t:
                pages.append(f"[Page {i}]\n{t}")
            if sum(len(p) for p in pages) > MAX_DOC_CHARS:
                break
        text = "\n\n".join(pages)
        if not text.strip():
            raise UnsupportedFile("the PDF has no text layer (scanned images are not read)")
    elif ext == ".docx":
        import docx

        d = docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        text = "\n".join(parts)
    elif ext in TEXT_EXT or _looks_like_text(data):
        text = _decode(data)
    else:
        raise UnsupportedFile(f"{ext or 'this file type'} is not a supported document")
    text = text.replace("\r\n", "\n").strip()
    truncated = len(text) > MAX_DOC_CHARS
    return text[:MAX_DOC_CHARS], truncated


def _decode(data):
    for enc in ("utf-8-sig", "utf-16") if data[:2] in (b"\xff\xfe", b"\xfe\xff") else ("utf-8-sig", "cp1252"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", "replace")


def _looks_like_text(data):
    sample = data[:4096]
    return bool(sample) and b"\x00" not in sample


def split_long_audio(audio, sr=16000, target_s=20.0, window_s=5.0):
    """Cuts long audio into ~20 s pieces at the quietest 100 ms near each boundary."""
    n = len(audio)
    step = int(target_s * sr)
    if n <= step * 1.25:
        return [audio]
    hop = sr // 10
    pieces, start = [], 0
    while n - start > step * 1.25:
        lo = start + step - int(window_s * sr)
        hi = min(n - hop, start + step + int(window_s * sr))
        frames = np.lib.stride_tricks.sliding_window_view(audio[lo:hi], hop)[::hop // 2]
        cut = lo + int(np.argmin(np.sqrt((frames ** 2).mean(axis=1)))) * (hop // 2) + hop // 2
        pieces.append(audio[start:cut])
        start = cut
    pieces.append(audio[start:])
    return pieces


def _duration(sec):
    sec = int(round(sec or 0))
    return f"{sec // 60}:{sec % 60:02d}"


def public(att):
    """Attachment metadata for the browser (no big payloads except image thumbnails)."""
    out = {k: att.get(k) for k in ("id", "kind", "name", "duration", "truncated")}
    if att.get("kind") == "image":
        out["data"] = att.get("data")
    elif att.get("kind") == "video" and att.get("frames"):
        out["data"] = att["frames"][0]
    return out


def llm_content(text, attachments, vision):
    """Builds the user message content (string, or parts list when images are sent)."""
    blocks, images, notes = [], [], []
    for a in attachments or []:
        kind, name = a.get("kind"), a.get("name", "file")
        if kind == "document":
            cut = " (truncated)" if a.get("truncated") else ""
            blocks.append(f"[Attached file: {name}{cut}]\n{a.get('text', '')}\n[End of {name}]")
        elif kind == "audio":
            blocks.append(f"[Attached audio: {name}, {_duration(a.get('duration'))}. Transcript:]\n"
                          f"{a.get('text') or '(no speech detected)'}\n[End of transcript]")
        elif kind == "video":
            head = f"[Attached video: {name}, {_duration(a.get('duration'))}"
            if a.get("frames") and vision:
                head += f", {len(a['frames'])} frames shown in order below"
                images += a["frames"]
            blocks.append(head + ". Audio transcript:]\n" + (a.get("text") or "(no speech)") + "\n[End of transcript]")
        elif kind == "image":
            if vision and a.get("data"):
                images.append(a["data"])
            else:
                notes.append(f"[An image '{name}' was attached but the current model cannot see images.]")
    body = "\n\n".join([t for t in [text] + blocks + notes if t])
    if not images:
        return body
    return [{"type": "text", "text": body}] + [{"type": "image_url", "image_url": {"url": u}} for u in images]
